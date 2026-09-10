"""方案文档自动生成与导出（指导书 6.7）。

模板+内容装配：章节骨架来自模板库，表格与数值由选型/测算服务直出；
python-docx 渲染 Word，LibreOffice 无头模式转 PDF（未安装时仅 Word）。
"""

from __future__ import annotations

import logging
import shutil
import subprocess
import uuid
from datetime import datetime, UTC
from io import BytesIO

from docx import Document
from docx.shared import Pt

from backend.app.core.config import settings

logger = logging.getLogger("icops.documents")
AI_DISCLAIMER = "AI 生成初稿，需人工确认"


def _soffice() -> str | None:
    """探测 LibreOffice 可执行文件：环境变量 -> PATH -> Windows 常见安装目录。"""
    cfg = settings.libreoffice_path
    if cfg and cfg != "auto":
        return cfg
    import os

    candidates: list[str] = []
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            candidates.append(p)
    if os.name == "nt":
        roots = [
            os.environ.get("ProgramFiles", "C:\\Program Files"),
            os.environ.get("ProgramFiles(x86)", "C:\\Program Files (x86)"),
            os.environ.get("LOCALAPPDATA", ""),
        ]
        for r in roots:
            if not r:
                continue
            candidates.append(os.path.join(r, "LibreOffice", "program", "soffice.exe"))
            candidates.append(os.path.join(r, "libreoffice", "program", "soffice.exe"))
    for c in candidates:
        if c and os.path.isfile(c):
            return c
    return None


def _head(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run.font.size = Pt(18 if level == 0 else 14 if level == 1 else 12)


def _para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "宋体"
    return p


def render_docx(payload: dict, doc_type: str, title: str) -> bytes:
    """按 doc_type 渲染 Word 文档字节流。payload 由方案生成智能体产出。"""
    doc = Document()
    _head(doc, title, 0)
    meta = payload.get("meta", {})
    _para(
        doc,
        f"生成时间：{datetime.now(UTC).strftime('%Y-%m-%d %H:%M')}　"
        f"模型版本：{meta.get('model_version', 'demo-template')}　知识版本：{meta.get('kb_version', 'V1.0')}",
    )
    doc.add_paragraph()

    if doc_type in ("selection", "bid", "construction"):
        # 需求描述
        req = payload.get("requirement", {})
        if req:
            _head(doc, "一、客户需求与工况", 1)
            _para(
                doc,
                f"场景：{req.get('scene_cn', req.get('scene_type', ''))}；"
                f"年产能：{req.get('annual_t', 0):,.0f} 吨；"
                f"预算：{req.get('budget_cny', 0):,.0f} 元；"
                f"约束：{('、'.join(req.get('constraints') or [])) or '无'}",
            )
        bundles = payload.get("bundles") or []
        if bundles:
            _head(doc, "二、设备选型方案（Top3）", 1)
            for i, b in enumerate(bundles, 1):
                _head(doc, f"2.{i} {b.get('name', f'方案{i}')}", 2)
                _para(doc, b.get("summary", ""))
                tbl = doc.add_table(rows=1, cols=4)
                tbl.style = "Light Grid Accent 1"
                for j, c in enumerate(["型号", "类别", "数量", "单价(元)"]):
                    tbl.rows[0].cells[j].text = c
                for o in b.get("fleet", []):
                    row = tbl.add_row().cells
                    row[0].text = f"{o.get('model_code')} {o.get('model_name', '')}"
                    row[1].text = o.get("category_cn", "")
                    row[2].text = str(o.get("count"))
                    row[3].text = f"{o.get('unit_price_cny', 0):,.0f}"
                _head(doc, "三年 TCO（四大类：购置/能耗/维保/残值）", 3)
                t2 = doc.add_table(rows=1, cols=6)
                for j, c in enumerate(["型号", "数量", "购置(元)", "能耗3年(元)", "维保3年(元)", "残值(元)"]):
                    t2.rows[0].cells[j].text = c
                for tc in b.get("tco", []):
                    cells = t2.add_row().cells
                    cells[0].text = tc.get("model_code", "")
                    cells[1].text = str(tc.get("count", 1))
                    cells[2].text = f"{tc.get('purchase_cny', 0):,.0f}"
                    cells[3].text = f"{tc.get('energy_3y_cny', 0):,.0f}"
                    cells[4].text = f"{tc.get('maintenance_3y_cny', 0):,.0f}"
                    cells[5].text = f"{tc.get('residual_cny', 0):,.0f}"
        # 通用章节内容（模板章节 + AI 生成段落）
        for i, ch in enumerate(payload.get("chapters") or [], 3):
            _head(doc, f"{i}、{ch.get('chapter', '')}", 1)
            for line in ch.get("paragraphs") or ch.get("bullets") or []:
                _para(doc, f"· {line}")
        if doc_type == "bid":
            resp = payload.get("bid_response_check")
            if resp:
                _head(doc, "招标响应度检查", 1)
                for item in resp:
                    _para(doc, f"[{item.get('status', '')}] {item.get('clause', '')}")
    else:
        for ch in payload.get("chapters") or []:
            _head(doc, ch.get("chapter", ""), 1)
            for line in ch.get("paragraphs") or []:
                _para(doc, line)

    _para(doc, "")
    _para(
        doc, f"※ {AI_DISCLAIMER}。本方案数值均由结构化数据库直读，可逐项溯源；正式商务文件请人工复核后签发。"
    )
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_document(payload: dict, doc_type: str, title: str) -> dict:
    """渲染并保存 docx；如环境有 LibreOffice 再转 PDF。返回文件路径信息。"""
    out_dir = settings.repo_root / "data" / "generated"
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(c for c in title if c not in '\\/:*?"<>|')[:40] or "方案"
    docx_bytes = render_docx(payload, doc_type, title)
    docx_path = out_dir / f"{safe}_{uuid.uuid4().hex[:6]}.docx"
    docx_path.write_bytes(docx_bytes)
    pdf_path = None
    soffice = _soffice()
    if soffice:
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                check=True,
                capture_output=True,
                timeout=120,
            )
            cand = docx_path.with_suffix(".pdf")
            if cand.exists():
                pdf_path = str(cand)
        except Exception as exc:  # noqa: BLE001
            logger.warning("LibreOffice PDF 转换失败（仅提供 Word）：%s", exc)
    return {"docx": str(docx_path), "pdf": pdf_path}
